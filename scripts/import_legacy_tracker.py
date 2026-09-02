#!/usr/bin/env python3
"""
Extracts and cleans data from the firm's original "Trump Administration
Executive Actions Tracker" spreadsheet (maintained through Jan 20, 2026) into
JSON seed files matching this app's TypeScript types (src/lib/types.ts).

Usage:
    python3 scripts/import_legacy_tracker.py

Reads:  data/source/trump-admin-executive-actions-tracker.xlsx
Writes: src/data/legacy-import/executive-orders.json
        src/data/legacy-import/rescinded-prior-orders.json
        src/data/legacy-import/agency-actions.json

Re-run this whenever the source spreadsheet is updated with a fresher export
from the firm's Google Sheet. The app reads these JSON files as its local
fallback dataset (see src/lib/data.ts) until Supabase is connected — once it
is, use `npm run import:supabase` to load the same JSON into the database.

Requires: openpyxl (pip install openpyxl)
"""

import json
import re
from datetime import datetime
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parent.parent
SOURCE_XLSX = REPO_ROOT / "data" / "source" / "trump-admin-executive-actions-tracker.xlsx"
OUT_DIR = REPO_ROOT / "src" / "data" / "legacy-import"

# Fields the firm hand-curated (not AI-generated) — flagged so a future
# automated enrichment pass doesn't overwrite them without review.
CURATED_FIELDS = [
    "subjectArea",
    "aiSummary",
    "deliverable",
    "timelineNotes",
    "availableAnalysis",
    "legalChallenges",
]


def iso_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    return None


def clean_text(value):
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def split_multiline_list(value):
    """Splits a newline- (or wide-gap-) separated cell into a list of trimmed strings."""
    if not value:
        return []
    text = str(value)
    parts = [p.strip() for p in text.split("\n")]
    parts = [p for p in parts if p]
    # A few rows use runs of spaces instead of newlines (e.g. "DHS     DOS").
    expanded = []
    for p in parts:
        if re.search(r"\s{3,}", p):
            expanded.extend([x.strip() for x in re.split(r"\s{3,}", p) if x.strip()])
        else:
            expanded.append(p)
    return expanded


def parse_type_number(raw):
    """Splits the source "Type/Number" column into (eo_number, action_type)."""
    if raw is None:
        return None, "Unspecified"

    # A few rows store the EO number as a bare number (e.g. 14314.0) with no
    # "EO" prefix at all.
    if isinstance(raw, (int, float)):
        return f"EO {int(raw)}", "Executive Order"

    text = str(raw).strip()
    if not text:
        return None, "Unspecified"

    if re.match(r"^\d+$", text):
        return f"EO {text}", "Executive Order"

    m = re.match(r"^(?:EO|E\.O\.|Executive Order)\.?\s*(\d+)", text, re.IGNORECASE)
    if m:
        return f"EO {m.group(1)}", "Executive Order"

    if re.match(r"^(?:EO|E\.O\.|Executive Order)\b", text, re.IGNORECASE):
        # e.g. "EO (not yet listed)" — a signed EO with no number assigned yet.
        return None, "Executive Order"

    if re.match(r"^Proclamation\b", text, re.IGNORECASE):
        return None, text

    if "memo" in text.lower() or "nspm" in text.lower():
        return None, "Memorandum"

    if "not posted" in text.lower():
        return None, "Pending Federal Register Publication"

    return None, text


LEGAL_CHALLENGE_PATTERN = re.compile(
    r"^(?P<case>.+?)\s*\(\s*(?P<court>[^)]+?)\s*\)\s*(?:[-–—]\s*(?P<detail>.*))?$",
    re.DOTALL,
)


def parse_legal_challenges(raw):
    if not raw or str(raw).strip().lower() in ("none", "n/a", ""):
        return []
    entries = re.split(r"\n\s*\n", str(raw).strip())
    challenges = []
    for entry in entries:
        entry = entry.strip()
        if not entry:
            continue
        m = LEGAL_CHALLENGE_PATTERN.match(entry)
        if m:
            case = m.group("case").strip()
            court = m.group("court").strip()
            detail = (m.group("detail") or "").strip()
            challenges.append(
                {
                    "caseName": case,
                    "court": court,
                    "status": detail if detail else "See summary",
                    "summary": detail if detail else entry,
                }
            )
        else:
            challenges.append(
                {
                    "caseName": "Legal challenge",
                    "court": "",
                    "status": "See summary",
                    "summary": entry,
                }
            )
    return challenges


def extract_executive_orders(wb):
    ws = wb["Trump Admin Exec Actions"]
    results = []
    idx = 0
    for r in range(4, ws.max_row + 1):
        title = clean_text(ws.cell(row=r, column=2).value)
        key_date = ws.cell(row=r, column=3).value
        type_number_raw = ws.cell(row=r, column=4).value
        type_number = clean_text(type_number_raw)
        if not title:
            continue
        # Rows like "Litigation Tracker" / "Media Trackers" are reference
        # links, not executive actions — they carry neither a date nor a
        # type/number.
        if key_date is None and type_number is None:
            continue

        idx += 1
        eo_number, action_type = parse_type_number(type_number_raw)
        subject_area = clean_text(ws.cell(row=r, column=1).value)

        now = datetime.utcnow().isoformat() + "Z"
        results.append(
            {
                "id": f"legacy-eo-{idx}",
                "eoNumber": eo_number,
                "actionType": action_type,
                "title": title,
                "federalRegisterUrl": None,
                "dateSigned": iso_date(key_date) or "",
                "datePublished": None,
                "status": "active",
                "agenciesImpacted": split_multiline_list(ws.cell(row=r, column=5).value),
                "keyDates": [],
                "subjectArea": [subject_area] if subject_area else [],
                "practiceAreas": [],
                "industries": [],
                "aiSummary": clean_text(ws.cell(row=r, column=6).value),
                "timelineNotes": clean_text(ws.cell(row=r, column=7).value),
                "deliverable": clean_text(ws.cell(row=r, column=8).value),
                "availableAnalysis": clean_text(ws.cell(row=r, column=9).value),
                "legalChallenges": parse_legal_challenges(ws.cell(row=r, column=10).value),
                "newsMentions": [],
                "manuallyEditedFields": CURATED_FIELDS,
                "createdAt": now,
                "updatedAt": now,
            }
        )
    return results


def extract_rescinded_prior_orders(wb):
    ws = wb["Rescinded Exec Actions"]
    results = []
    idx = 0
    for r in range(4, ws.max_row + 1):
        title = clean_text(ws.cell(row=r, column=3).value)
        if not title:
            continue
        idx += 1
        order_number_raw = ws.cell(row=r, column=1).value
        order_number = None
        if order_number_raw is not None:
            # Stored as a float in the sheet (e.g. 11246.0) — render as a plain int string.
            try:
                order_number = str(int(order_number_raw))
            except (TypeError, ValueError):
                order_number = clean_text(order_number_raw)

        now = datetime.utcnow().isoformat() + "Z"
        results.append(
            {
                "id": f"legacy-rescinded-{idx}",
                "orderNumber": order_number,
                "dateSigned": iso_date(ws.cell(row=r, column=2).value),
                "title": title,
                "administration": clean_text(ws.cell(row=r, column=4).value) or "Unknown",
                "createdAt": now,
            }
        )
    return results


def extract_agency_actions(wb):
    ws = wb["Select Agency Actions"]
    results = []
    idx = 0
    for r in range(5, ws.max_row + 1):
        title = clean_text(ws.cell(row=r, column=1).value)
        if not title:
            continue
        idx += 1
        now = datetime.utcnow().isoformat() + "Z"
        results.append(
            {
                "id": f"legacy-agency-action-{idx}",
                "title": title,
                "issuingAgency": clean_text(ws.cell(row=r, column=2).value) or "Unspecified",
                "keyDate": iso_date(ws.cell(row=r, column=3).value),
                "otherAgenciesImpacted": split_multiline_list(ws.cell(row=r, column=4).value),
                "legalChallenges": parse_legal_challenges(ws.cell(row=r, column=5).value),
                "availableAnalysis": clean_text(ws.cell(row=r, column=6).value),
                "relatedEoNumber": clean_text(ws.cell(row=r, column=7).value),
                "createdAt": now,
            }
        )
    return results


def main():
    if not SOURCE_XLSX.exists():
        raise SystemExit(f"Source spreadsheet not found at {SOURCE_XLSX}")

    wb = openpyxl.load_workbook(SOURCE_XLSX, data_only=True)
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    datasets = {
        "executive-orders.json": extract_executive_orders(wb),
        "rescinded-prior-orders.json": extract_rescinded_prior_orders(wb),
        "agency-actions.json": extract_agency_actions(wb),
    }

    for filename, data in datasets.items():
        out_path = OUT_DIR / filename
        with out_path.open("w", encoding="utf-8") as f:
            json.dump(data, f, indent=2, ensure_ascii=False)
        print(f"Wrote {len(data)} records to {out_path.relative_to(REPO_ROOT)}")


if __name__ == "__main__":
    main()
